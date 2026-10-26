"""Optional Office export adapters: page images to slides and visible DOCX signature.
DOCX signature preserves package parts and adds a floating picture on the final section;
no cryptographic signing or OCR is represented by these operations.
"""
import io,zipfile,math,xml.etree.ElementTree as ET
from pathlib import Path
import engine as base

def pdf_pptx(job):
 from pptx import Presentation
 from pptx.util import Emu
 prs=Presentation();count=0;dpi=int(job.get('dpi',144))
 if not 72<=dpi<=300:raise ValueError('invalidRequest')
 with base.open_pdf(job['path']) as doc:
  if doc.page_count>500:raise ValueError('fileTooLarge')
  first=doc[0].rect;prs.slide_width=Emu(round(first.width*12700));prs.slide_height=Emu(round(first.height*12700))
  for page in doc:
   width=math.ceil(page.rect.width*dpi/72);height=math.ceil(page.rect.height*dpi/72)
   if width<1 or height<1 or width*height>32_000_000:raise ValueError('fileTooLarge')
   pix=page.get_pixmap(dpi=dpi,alpha=False)
   if pix.width*pix.height>80_000_000:raise ValueError('fileTooLarge')
   stream=io.BytesIO(pix.tobytes('png'));slide=prs.slides.add_slide(prs.slide_layouts[6]);ratio=min(prs.slide_width/(page.rect.width*12700),prs.slide_height/(page.rect.height*12700));w=round(page.rect.width*12700*ratio);h=round(page.rect.height*12700*ratio)
   slide.shapes.add_picture(stream,Emu((prs.slide_width-w)//2),Emu((prs.slide_height-h)//2),Emu(w),Emu(h));count+=1
 out=base.write_atomic(job['output_dir'],Path(job['path']).stem+'_slides','.pptx',lambda p:prs.save(str(p)))
 return base.result([out],{'mode':'pageImages','pages':count})

def docx_signature(path,signature,job):
 NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main','pic':'http://schemas.openxmlformats.org/drawingml/2006/picture','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
 with zipfile.ZipFile(path) as z:
  if sum(i.file_size for i in z.infolist())>100_000_000:raise ValueError('fileTooLarge')
  parts={i.filename:z.read(i) for i in z.infolist()}
 document=parts.get('word/document.xml')
 if not document or b'<!DOCTYPE'in document or b'<!ENTITY'in document:raise ValueError('unsupported')
 root=ET.fromstring(document);body=root.find('{'+NS['w']+'}body')
 if body is None:raise ValueError('unsupported')
 x=int(float(job.get('x',20))*12700);y=int(float(job.get('y',20))*12700);w=signature.width*9525;h=signature.height*9525
 rid='rIdADMINSignature';name='admin-signature.png';mediaNumber=1
 while 'word/media/'+name in parts:mediaNumber+=1;name=f'admin-signature-{mediaNumber}.png'
 docId=max([int(x.get('id','0')) for x in root.iter('{'+NS['wp']+'}docPr') if str(x.get('id','0')).isdigit()]+[10000])+1
 relationships=parts.get('word/_rels/document.xml.rels',b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
 rel=ET.fromstring(relationships)
 n=1;existing={e.get('Id') for e in rel}
 while rid in existing:n+=1;rid='rIdADMINSignature'+str(n)
 ET.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',{'Id':rid,'Type':NS['r']+'/image','Target':'media/'+name})
 floating=f'''<w:p xmlns:w="{NS['w']}" xmlns:wp="{NS['wp']}" xmlns:a="{NS['a']}" xmlns:pic="{NS['pic']}" xmlns:r="{NS['r']}"><w:r><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH><wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV><wp:extent cx="{w}" cy="{h}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="{docId}" name="ADMIN signature"/><wp:cNvGraphicFramePr/><a:graphic><a:graphicData uri="{NS['pic']}"><pic:pic><pic:nvPicPr><pic:cNvPr id="0" name="Signature"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r></w:p>'''
 paragraph=ET.fromstring(floating);section=body.find('{'+NS['w']+'}sectPr')
 if section is not None:body.insert(list(body).index(section),paragraph)
 else:body.append(paragraph)
 for prefix,ns in NS.items():ET.register_namespace(prefix,ns)
 parts['word/document.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True);parts['word/_rels/document.xml.rels']=ET.tostring(rel,encoding='utf-8',xml_declaration=True)
 png=io.BytesIO();signature.save(png,'PNG');parts['word/media/'+name]=png.getvalue()
 content=ET.fromstring(parts['[Content_Types].xml']);ns='http://schemas.openxmlformats.org/package/2006/content-types'
 if not any(e.get('Extension')=='png' for e in content):ET.SubElement(content,'{'+ns+'}Default',{'Extension':'png','ContentType':'image/png'})
 parts['[Content_Types].xml']=ET.tostring(content,encoding='utf-8',xml_declaration=True)
 def save(p):
  with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as out:
   for name,data in parts.items():out.writestr(name,data)
 output=base.write_atomic(job['output_dir'],path.stem+'_signed','.docx',save)
 return base.result([output],{'mode':'visibleSignature','placement':'finalSection','notCertificateSignature':True})
